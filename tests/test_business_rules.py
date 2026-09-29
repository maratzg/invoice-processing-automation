from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.config import MoneoTemplateConfig
from src.exceptions import TemplateError
from src.invoice_models import ClientRecord, ContainerInfo, Invoice, InvoiceHeader, InvoiceLine, VATRecord
from src.mapping_loader import MappingData
from src.moneo_writer import MONEO_COLUMNS, build_moneo_rows, write_moneo_import
from src.pdf_parser import parse_invoice_text
from src.validator import validate_invoice


def _template_config() -> MoneoTemplateConfig:
    return MoneoTemplateConfig(
        template_sheet=None,
        header_row=None,
        data_start_row=None,
        invoice_type="Rēķins",
        invoice_type_value="REK",
        currency="EUR",
        discount_percent=None,
    )


def _invoice() -> Invoice:
    return Invoice(
        source_path=Path("invoice.pdf"),
        header=InvoiceHeader(
            invoice_number="INV-1",
            invoice_date=date(2026, 1, 15),
            client_name="Client A",
            client_vat_reg_number="LV123",
            due_date=date(2026, 1, 30),
            payment_terms="14 days",
            subtotal=Decimal("300.00"),
            vat_amount=Decimal("21.00"),
            total_amount_eur_incl_vat=Decimal("321.00"),
        ),
        lines=[
            InvoiceLine(service_code="O.F.", description="Ocean freight", vat_rate=Decimal("0"), line_amount_ex_vat=Decimal("200.00")),
            InvoiceLine(service_code="THC/D", description="THC", vat_rate=Decimal("21"), line_amount_ex_vat=Decimal("100.00")),
        ],
        containers=[
            ContainerInfo("ABCD1234567", unit_type="40HC"),
            ContainerInfo("EFGH1234567", unit_type="40HC"),
        ],
    )


def _mappings() -> MappingData:
    return MappingData(
        clients=[ClientRecord("Client A", "LV123", "MONEO Client A", "C001", "LV", "Domestic")],
        vats=[
            VATRecord("DEFAULT", "DEFAULT", Decimal("21"), "21 %", "21"),
            VATRecord("LV", "Domestic", Decimal("0"), "0 %", "0EXP"),
        ],
        articles=[],
    )


class BusinessRuleTests(unittest.TestCase):
    def test_valid_invoice_builds_grouped_moneo_rows(self) -> None:
        invoice = _invoice()
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)
        rows, warnings = build_moneo_rows(invoice, outcome, _template_config())

        self.assertTrue(outcome.valid, outcome.errors)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].quantity, 2)
        self.assertEqual(rows[0].line_sum_ex_vat, Decimal("200.00"))
        self.assertEqual(rows[0].unit_price_ex_vat, Decimal("100.00"))
        self.assertEqual(warnings, [])

    def test_invoice_without_containers_is_rejected(self) -> None:
        invoice = _invoice()
        invoice.containers = []
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)

        self.assertFalse(outcome.valid)
        self.assertTrue(any("container" in error.lower() for error in outcome.errors))

    def test_missing_zero_vat_code_requires_review(self) -> None:
        invoice = _invoice()
        mappings = _mappings()
        mappings.vats[1].pvn_kods = ""
        outcome = validate_invoice(invoice, mappings, set(), template_available=True)

        self.assertFalse(outcome.valid)
        self.assertTrue(any("manual review" in error.lower() or "missing pvn" in error.lower() for error in outcome.errors))

    def test_writer_uses_template_column_names(self) -> None:
        invoice = _invoice()
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)
        rows, _ = build_moneo_rows(invoice, outcome, _template_config())

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            template_path = tmp_path / "Rekinu_Imp_2024.xlsx"
            output_path = tmp_path / "moneo_import_ready.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Import"
            sheet.append(MONEO_COLUMNS)
            workbook.save(template_path)

            write_moneo_import(template_path, output_path, rows, _template_config())
            written = load_workbook(output_path, data_only=True)
            sheet = written.active

            self.assertEqual(sheet["A2"].value.date(), date(2026, 1, 15))
            self.assertEqual(sheet["B2"].value, "MONEO Client A")
            self.assertEqual(sheet["L2"].value, 2)
            self.assertEqual(sheet["N2"].value, Decimal("200.00"))

    def test_parser_handles_shipping_table_lines_without_percent_signs(self) -> None:
        text = """
        Invoice INV-2026-001
        02.07.2026
        Page 1 / 2
        VAT Registration No. Your Reference Due Date
        LV40000000001 23.07.2026
        Project Code B/L No. Payment Terms
        DEMO2026-001 SYNTHBL000001 21 days
        Place of Receipt Port of Loading Port of Discharge Place of Delivery
        HAMBURG HAMBURG TALLINN TALLINN
        No. Description Quantity Excl. VAT VAT % Excl. VAT
        ISPS/D Int.Ship/Port Fac. Sec. Surch 2 1,00 10,00 10,00 0 20,00
        THC/D Terminal Handling Charge/Disc 2 1,00 140,00 140,00 0 280,00
        D/F Documentation Fee 2 1,00 130,00 130,00 21 260,00
        MERCH Merchant Haulage 2 1,00 45,00 45,00 21 90,00
        Containerno /Unit /Weight /CBM
        ABCD1234567 / 4SH / 25 000 / 40
        WXYZ7654321 / 4SH / 24 500 / 39,500
        Subtotal 650,00
        VAT Amount 73,50
        Total EUR Incl. VAT 723,50
        """
        invoice = parse_invoice_text(text, Path("sample.pdf"))

        self.assertEqual(invoice.header.invoice_number, "INV-2026-001")
        self.assertEqual(invoice.header.invoice_date, date(2026, 7, 2))
        self.assertEqual(invoice.header.due_date, date(2026, 7, 23))
        self.assertEqual(invoice.header.payment_terms, "21 days")
        self.assertEqual([line.line_amount_ex_vat for line in invoice.lines], [
            Decimal("20.00"),
            Decimal("280.00"),
            Decimal("260.00"),
            Decimal("90.00"),
        ])
        self.assertEqual([line.vat_rate for line in invoice.lines], [
            Decimal("0"),
            Decimal("0"),
            Decimal("21"),
            Decimal("21"),
        ])
        self.assertEqual(invoice.containers[0].weight, Decimal("25000"))
        self.assertEqual(invoice.containers[1].cbm, Decimal("39.500"))

    def test_validation_rejects_negative_line_amount(self) -> None:
        invoice = _invoice()
        invoice.lines[0].line_amount_ex_vat = Decimal("-1.00")
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)

        self.assertFalse(outcome.valid)
        self.assertTrue(any("negative amount" in error.lower() for error in outcome.errors))

    def test_validation_rejects_due_date_before_invoice_date(self) -> None:
        invoice = _invoice()
        invoice.header.due_date = date(2026, 1, 1)
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)

        self.assertFalse(outcome.valid)
        self.assertTrue(any("before invoice date" in error.lower() for error in outcome.errors))

    def test_validation_rejects_duplicate_invoice_in_current_batch(self) -> None:
        invoice = _invoice()
        outcome = validate_invoice(
            invoice,
            _mappings(),
            set(),
            template_available=True,
            batch_invoice_numbers_seen={"INV-1"},
        )

        self.assertFalse(outcome.valid)
        self.assertTrue(any("current pdf batch" in error.lower() for error in outcome.errors))

    def test_validation_allows_missing_pdf_client_name_when_vat_mapping_is_exact(self) -> None:
        invoice = _invoice()
        invoice.header.client_name = None
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)

        self.assertTrue(outcome.valid, outcome.errors)
        self.assertTrue(any("client name was not extracted" in warning.lower() for warning in outcome.warnings))

    def test_writer_preserves_template_example_row_and_formula_columns(self) -> None:
        invoice = _invoice()
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)
        rows, _ = build_moneo_rows(invoice, outcome, _template_config())

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            template_path = tmp_path / "Rekinu_Imp_2024.xlsx"
            output_path = tmp_path / "moneo_import_ready.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(list(range(1, 19)))
            sheet.append(["Obligāts"] * 18)
            sheet.append(
                [
                    1,
                    "Rēķina datums",
                    "Klienta nosaukums\n(kontakta kartiņai jābūt sistēmā)",
                    "Klienta Reģ.nr.\n(kontakta kartiņai jābūt sistēmā)",
                    "Rēķina veids",
                    "Rēķina veids value (nemainīt)",
                    "Valūta",
                    "Samaksas termiņš (dienas)",
                    "Pavadzīmes / Rēķina nr",
                    "Artikula kods (ja tie ir sistēmā)",
                    "Artikula nosaukums (rindas nosaukums)",
                    "Cena BEZ PVN",
                    "Daudzums",
                    "Atlaides % (kā skaitlis)",
                    "Rindas summa BEZ PVN",
                    "PVN likme",
                    "PVN kods (nemainīt)",
                    None,
                ]
            )
            sheet.append(
                [
                    None,
                    "01.01.2021",
                    "Example",
                    "LV0",
                    "Rēķins",
                    '=IF(E4>0,VLOOKUP(E4,Set!$A$2:$B$6,2,FALSE),"")',
                    "EUR",
                    "10",
                    "ML 210001",
                    None,
                    "Example row",
                    1,
                    1,
                    None,
                    1,
                    "21 %",
                    '=IF(P4>0,VLOOKUP(P4,Set!$H$2:$J$19,2,FALSE),"")',
                    None,
                ]
            )
            sheet.append([None, None, None, None, None, '=IF(E5>0,VLOOKUP(E5,Set!$A$2:$B$6,2,FALSE),"")', None, None, None, None, None, None, None, None, None, None, '=IF(P5>0,VLOOKUP(P5,Set!$H$2:$J$19,2,FALSE),"")', None])
            set_sheet = workbook.create_sheet("Set")
            set_sheet["A2"] = "Rēķins"
            set_sheet["B2"] = "REK"
            set_sheet["H2"] = "21 %"
            set_sheet["I2"] = 21
            set_sheet["H3"] = "0 %"
            set_sheet["I3"] = "0EXP"
            workbook.save(template_path)

            write_moneo_import(template_path, output_path, rows, _template_config())
            written = load_workbook(output_path, data_only=False)
            sheet = written["Sheet1"]

            self.assertEqual(sheet["B4"].value, "01.01.2021")
            self.assertEqual(sheet["B5"].value.date(), date(2026, 1, 15))
            self.assertTrue(str(sheet["F5"].value).startswith("="))
            self.assertTrue(str(sheet["Q5"].value).startswith("="))

    def test_writer_rejects_template_vat_code_mismatch(self) -> None:
        invoice = _invoice()
        outcome = validate_invoice(invoice, _mappings(), set(), template_available=True)
        rows, _ = build_moneo_rows(invoice, outcome, _template_config())
        rows[0].pvn_kods = "WRONG"

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            template_path = tmp_path / "Rekinu_Imp_2024.xlsx"
            output_path = tmp_path / "moneo_import_ready.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Import"
            sheet.append(MONEO_COLUMNS)
            set_sheet = workbook.create_sheet("Set")
            set_sheet["A2"] = "Rēķins"
            set_sheet["B2"] = "REK"
            set_sheet["H2"] = "0 %"
            set_sheet["I2"] = "0EXP"
            set_sheet["H3"] = "21 %"
            set_sheet["I3"] = 21
            workbook.save(template_path)

            with self.assertRaises(TemplateError):
                write_moneo_import(template_path, output_path, rows, _template_config())


if __name__ == "__main__":
    unittest.main()
