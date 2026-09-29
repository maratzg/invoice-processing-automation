from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .exceptions import MappingError
from .invoice_models import ArticleRecord, ClientRecord, VATRecord
from .pdf_parser import parse_money


CLIENT_FILE = "Client_Mapping.xlsx"
VAT_FILE = "VAT_Mapping.xlsx"
ARTICLE_FILE = "Article_Mapping.xlsx"
CLIENT_HEADERS = [
    "PDF client name",
    "PDF VAT/registration number",
    "MONEO client name",
    "MONEO client code",
    "Country",
    "Client type",
]
VAT_HEADERS = ["Country", "Client type", "VAT %", "PVN likme", "PVN kods", "Manual review required"]
ARTICLE_HEADERS = ["VAT %", "Artikula nosaukums", "Artikula kods"]


@dataclass(slots=True)
class MappingData:
    clients: list[ClientRecord]
    vats: list[VATRecord]
    articles: list[ArticleRecord]

    def find_client(self, pdf_client_name: str | None, vat_reg_number: str | None) -> ClientRecord | None:
        normalized_vat = _normalize_key(vat_reg_number)
        if normalized_vat:
            for client in self.clients:
                if _normalize_key(client.pdf_vat_reg_number) == normalized_vat:
                    return client

        normalized_name = _normalize_key(pdf_client_name)
        if normalized_name:
            for client in self.clients:
                if _normalize_key(client.pdf_client_name) == normalized_name:
                    return client
        return None

    def find_article(self, vat_rate: Decimal) -> ArticleRecord | None:
        for article in self.articles:
            if article.vat_rate == vat_rate:
                return article
        if vat_rate == Decimal("21"):
            return ArticleRecord(vat_rate=Decimal("21"), artikula_nosaukums="Pakalpojums")
        if vat_rate == Decimal("0"):
            return ArticleRecord(vat_rate=Decimal("0"), artikula_nosaukums="Pakalpojums ar PVN neapl.")
        return None

    def find_vat(self, vat_rate: Decimal, client: ClientRecord | None) -> VATRecord | None:
        country = _normalize_key(client.country if client else None)
        client_type = _normalize_key(client.client_type if client else None)

        candidates = [vat for vat in self.vats if vat.vat_rate == vat_rate]
        ranked: list[tuple[int, VATRecord]] = []
        for vat in candidates:
            score = 0
            record_country = _normalize_key(vat.country)
            record_type = _normalize_key(vat.client_type)
            if record_country and country and record_country == country:
                score += 4
            elif record_country in {"", "DEFAULT", "*"}:
                score += 1
            elif record_country == "EU" and country and country not in {"LV", "LATVIA"}:
                score += 2

            if record_type and client_type and record_type == client_type:
                score += 4
            elif record_type in {"", "DEFAULT", "*"}:
                score += 1

            if score:
                ranked.append((score, vat))

        if ranked:
            ranked.sort(key=lambda item: item[0], reverse=True)
            return ranked[0][1]

        if vat_rate == Decimal("21"):
            return VATRecord(country="DEFAULT", client_type="DEFAULT", vat_rate=Decimal("21"), pvn_likme="21 %", pvn_kods="21")
        return None


def load_mappings(mappings_dir: Path) -> MappingData:
    mapping_data = MappingData(
        clients=_load_clients(mappings_dir / CLIENT_FILE),
        vats=_load_vats(mappings_dir / VAT_FILE),
        articles=_load_articles(mappings_dir / ARTICLE_FILE),
    )
    _validate_mapping_data(mapping_data)
    return mapping_data


def create_mapping_templates(mappings_dir: Path, overwrite: bool = False) -> list[Path]:
    mappings_dir.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []

    created.append(
        _create_workbook_if_needed(
            mappings_dir / CLIENT_FILE,
            [
                "PDF client name",
                "PDF VAT/registration number",
                "MONEO client name",
                "MONEO client code",
                "Country",
                "Client type",
                "Notes",
            ],
            [
                [
                    "Example Client SIA",
                    "LV40000000000",
                    "Example Client SIA",
                    "C0001",
                    "LV",
                    "Domestic",
                    "Replace examples with real MONEO client mappings.",
                ]
            ],
            overwrite,
        )
    )
    created.append(
        _create_workbook_if_needed(
            mappings_dir / VAT_FILE,
            [
                "Country",
                "Client type",
                "VAT %",
                "PVN likme",
                "PVN kods",
                "Manual review required",
                "Notes",
            ],
            [
                ["DEFAULT", "DEFAULT", 21, "21 %", "21", "No", "Default local 21% VAT rule."],
                ["EU", "EU business", 0, "0 %", "", "Yes", "Fill the correct MONEO 0% VAT code before production use."],
                ["NON-EU", "Export", 0, "0 %", "", "Yes", "Fill the correct MONEO 0% VAT code before production use."],
            ],
            overwrite,
        )
    )
    created.append(
        _create_workbook_if_needed(
            mappings_dir / ARTICLE_FILE,
            ["VAT %", "Artikula nosaukums", "Artikula kods", "Notes"],
            [
                [21, "Pakalpojums", "", "Default 21% service article."],
                [0, "Pakalpojums ar PVN neapl.", "", "Default 0% service article."],
            ],
            overwrite,
        )
    )
    return created


def _load_clients(path: Path) -> list[ClientRecord]:
    rows = _read_rows(path, CLIENT_HEADERS)
    clients: list[ClientRecord] = []
    for row in rows:
        pdf_name = _string(row.get("PDF client name"))
        moneo_name = _string(row.get("MONEO client name"))
        if not pdf_name or not moneo_name or pdf_name.lower().startswith("example "):
            continue
        clients.append(
            ClientRecord(
                pdf_client_name=pdf_name,
                pdf_vat_reg_number=_string(row.get("PDF VAT/registration number")),
                moneo_client_name=moneo_name,
                moneo_client_code=_string(row.get("MONEO client code")),
                country=_string(row.get("Country")),
                client_type=_string(row.get("Client type")),
            )
        )
    return clients


def _load_vats(path: Path) -> list[VATRecord]:
    rows = _read_rows(path, VAT_HEADERS)
    vats: list[VATRecord] = []
    for row in rows:
        vat_rate = parse_money(row.get("VAT %"))
        if vat_rate is None:
            continue
        manual_review = _string(row.get("Manual review required")).lower() in {"yes", "true", "1", "y"}
        vats.append(
            VATRecord(
                country=_string(row.get("Country")),
                client_type=_string(row.get("Client type")),
                vat_rate=vat_rate,
                pvn_likme=_string(row.get("PVN likme")),
                pvn_kods=_string(row.get("PVN kods")),
                manual_review_required=manual_review,
            )
        )
    return vats


def _load_articles(path: Path) -> list[ArticleRecord]:
    rows = _read_rows(path, ARTICLE_HEADERS)
    articles: list[ArticleRecord] = []
    for row in rows:
        vat_rate = parse_money(row.get("VAT %"))
        article_name = _string(row.get("Artikula nosaukums"))
        if vat_rate is None or not article_name:
            continue
        articles.append(
            ArticleRecord(
                vat_rate=vat_rate,
                artikula_nosaukums=article_name,
                artikula_kods=_string(row.get("Artikula kods")),
            )
        )
    return articles


def _read_rows(path: Path, required_headers: list[str]) -> list[dict[str, Any]]:
    if not path.exists():
        raise MappingError(f"Required mapping workbook is missing: {path}")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise MappingError(f"Could not open mapping workbook {path.name}: {exc}") from exc
    sheet = workbook.active
    headers = [_string(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    header_lookup = {_normalize_header(header): header for header in headers if header}
    missing = [header for header in required_headers if _normalize_header(header) not in header_lookup]
    if missing:
        raise MappingError(f"{path.name} is missing required columns: {', '.join(missing)}")
    canonical_by_normalized = {_normalize_header(header): header for header in required_headers}
    canonical_headers = [canonical_by_normalized.get(_normalize_header(header), header) for header in headers]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append({canonical_headers[index]: value for index, value in enumerate(values) if index < len(canonical_headers)})
    return rows


def _validate_mapping_data(mapping_data: MappingData) -> None:
    client_keys: dict[str, ClientRecord] = {}
    for client in mapping_data.clients:
        for key in [_normalize_key(client.pdf_vat_reg_number), _normalize_key(client.pdf_client_name)]:
            if not key:
                continue
            existing = client_keys.get(key)
            if existing and existing.moneo_client_name != client.moneo_client_name:
                raise MappingError(
                    "Conflicting client mapping found for "
                    f"{client.pdf_client_name or client.pdf_vat_reg_number!r}: "
                    f"{existing.moneo_client_name!r} vs {client.moneo_client_name!r}"
                )
            client_keys[key] = client

    vat_keys: set[tuple[str, str, Decimal]] = set()
    for vat in mapping_data.vats:
        key = (_normalize_key(vat.country), _normalize_key(vat.client_type), vat.vat_rate)
        if key in vat_keys:
            raise MappingError(
                "Duplicate VAT mapping found for "
                f"Country={vat.country!r}, Client type={vat.client_type!r}, VAT={vat.vat_rate}%"
            )
        vat_keys.add(key)

    article_rates: set[Decimal] = set()
    for article in mapping_data.articles:
        if article.vat_rate in article_rates:
            raise MappingError(f"Duplicate article mapping found for VAT {article.vat_rate}%.")
        article_rates.add(article.vat_rate)


def _create_workbook_if_needed(path: Path, headers: list[str], rows: list[list[Any]], overwrite: bool) -> Path:
    if path.exists() and not overwrite:
        return path
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mapping"
    _write_table(sheet, headers, rows)
    workbook.save(path)
    return path


def _write_table(sheet: Worksheet, headers: list[str], rows: list[list[Any]]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row_index, row_values in enumerate(rows, start=2):
        for column_index, value in enumerate(row_values, start=1):
            sheet.cell(row=row_index, column=column_index, value=value)
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)


def _string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_key(value: str | None) -> str:
    if not value:
        return ""
    return "".join(str(value).upper().split())


def _normalize_header(value: str | None) -> str:
    return " ".join(str(value or "").strip().lower().split())
