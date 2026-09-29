from __future__ import annotations

import shutil
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .config import MoneoTemplateConfig
from .exceptions import AccountingRuleError, TemplateError
from .file_utils import atomic_save_workbook, atomic_write_path
from .invoice_models import Invoice, MoneoRow, ValidationOutcome, iso_datetime_now


MONEO_COLUMNS = [
    "Rēķina datums",
    "Klienta nosaukums",
    "Klienta Reģ.nr.",
    "Rēķina veids",
    "Rēķina veids value",
    "Valūta",
    "Samaksas termiņš",
    "Pavadzīmes / Rēķina nr",
    "Artikula kods",
    "Artikula nosaukums",
    "Cena BEZ PVN",
    "Daudzums",
    "Atlaides %",
    "Rindas summa BEZ PVN",
    "PVN likme",
    "PVN kods",
]
PROTECTED_FORMULA_COLUMNS = {"Rēķina veids value", "PVN kods"}


def build_moneo_rows(
    invoice: Invoice,
    outcome: ValidationOutcome,
    template_config: MoneoTemplateConfig,
) -> tuple[list[MoneoRow], list[str]]:
    if not outcome.valid or outcome.client is None:
        return [], []

    warnings: list[str] = []
    grouped = _group_amounts_by_vat(invoice)
    container_count = invoice.container_count
    if container_count <= 0:
        raise AccountingRuleError(f"Cannot build MONEO rows without containers: {invoice.source_filename}")
    if not invoice.header.invoice_date:
        raise AccountingRuleError(f"Cannot build MONEO rows without invoice date: {invoice.source_filename}")
    if not grouped:
        raise AccountingRuleError(f"Cannot build MONEO rows without grouped line amounts: {invoice.source_filename}")
    rows: list[MoneoRow] = []

    for vat_rate, grouped_amount in sorted(grouped.items()):
        vat_record = outcome.vat_records_by_rate[vat_rate]
        article_record = outcome.article_records_by_rate[vat_rate]
        line_sum = _round_money(grouped_amount)
        unit_price = _round_money(line_sum / Decimal(container_count))
        if _round_money(unit_price * Decimal(container_count)) != line_sum:
            warnings.append(
                f"Rounding warning for VAT {vat_rate}%: unit price x container count does not exactly equal line sum."
            )

        rows.append(
            MoneoRow(
                invoice_date=invoice.header.invoice_date,  # type: ignore[arg-type]
                client_name=outcome.client.moneo_client_name,
                client_reg_number=invoice.header.client_vat_reg_number or "",
                invoice_type=template_config.invoice_type,
                invoice_type_value=template_config.invoice_type_value,
                currency=template_config.currency,
                payment_terms=invoice.header.payment_terms or "",
                invoice_number=invoice.header.invoice_number or "",
                article_code=article_record.artikula_kods,
                article_name=article_record.artikula_nosaukums,
                unit_price_ex_vat=unit_price,
                quantity=container_count,
                discount_percent=template_config.discount_percent,
                line_sum_ex_vat=line_sum,
                pvn_likme=vat_record.pvn_likme or "",
                pvn_kods=vat_record.pvn_kods or "",
                source_pdf=invoice.source_filename,
                vat_rate=vat_rate,
            )
        )
    return rows, warnings


def write_moneo_import(
    template_path: Path,
    output_path: Path,
    rows: Iterable[MoneoRow],
    template_config: MoneoTemplateConfig,
) -> None:
    rows = list(rows)
    if not rows:
        raise TemplateError("Refusing to write MONEO import with zero rows.")
    if not template_path.exists():
        raise TemplateError(f"Official MONEO template is missing: {template_path}")
    if template_path.resolve() == output_path.resolve():
        raise TemplateError("Output path must not be the same file as the official MONEO template.")

    def _writer(temp_path: Path) -> None:
        shutil.copyfile(template_path, temp_path)
        workbook = load_workbook(temp_path)
        sheet = workbook[template_config.template_sheet] if template_config.template_sheet else workbook.active

        header_row = template_config.header_row or _find_header_row(sheet)
        column_map = _build_column_map(sheet, header_row)
        data_start_row = template_config.data_start_row or _find_data_start_row(sheet, column_map, header_row)
        _validate_template_lists(workbook, rows, template_config)
        _prepare_data_rows(sheet, column_map, data_start_row, len(rows))

        for row_offset, moneo_row in enumerate(rows):
            excel_row = data_start_row + row_offset
            for column_name, value in moneo_row.to_template_values().items():
                if column_name not in column_map:
                    continue
                cell = sheet.cell(row=excel_row, column=column_map[column_name])
                if column_name in PROTECTED_FORMULA_COLUMNS and _cell_has_formula(cell):
                    continue
                cell.value = value
                _format_template_cell(cell, column_name)

        _force_excel_recalculation(workbook)
        workbook.save(temp_path)

    atomic_write_path(output_path, _writer)


def write_no_import_workbook(output_path: Path, reason: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NO_IMPORT"
    sheet["A1"] = "MONEO import file was not created"
    sheet["A2"] = reason
    sheet["A4"] = "Place the official Rekinu_Imp_2024.xlsx template in templates/ and rerun after errors are fixed."
    sheet["A1"].font = Font(bold=True, size=14, color="9C0006")
    sheet.column_dimensions["A"].width = 120
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet["A4"].alignment = Alignment(wrap_text=True)
    atomic_save_workbook(workbook, output_path)


def write_extracted_data(output_path: Path, invoices: list[Invoice]) -> None:
    workbook = Workbook()
    header_sheet = workbook.active
    header_sheet.title = "Headers"
    _write_rows(
        header_sheet,
        [
            "Source PDF",
            "Invoice number",
            "Invoice date",
            "Client name",
            "Client VAT/registration",
            "Due date",
            "Payment terms",
            "Project code",
            "B/L number",
            "Place of receipt",
            "Port of loading",
            "Port of discharge",
            "Place of delivery",
            "Subtotal",
            "VAT amount",
            "Total incl VAT",
            "VAT clause",
            "Container count",
        ],
        [
            [
                invoice.source_filename,
                invoice.header.invoice_number,
                invoice.header.invoice_date,
                invoice.header.client_name,
                invoice.header.client_vat_reg_number,
                invoice.header.due_date,
                invoice.header.payment_terms,
                invoice.header.project_code,
                invoice.header.bl_number,
                invoice.header.place_of_receipt,
                invoice.header.port_of_loading,
                invoice.header.port_of_discharge,
                invoice.header.place_of_delivery,
                invoice.header.subtotal,
                invoice.header.vat_amount,
                invoice.header.total_amount_eur_incl_vat,
                invoice.header.vat_clause,
                invoice.container_count,
            ]
            for invoice in invoices
        ],
    )

    line_sheet = workbook.create_sheet("Lines")
    _write_rows(
        line_sheet,
        [
            "Source PDF",
            "Invoice number",
            "Service code",
            "Description",
            "Quantity",
            "Pricing currency",
            "Exchange rate",
            "Price in currency",
            "Currency code",
            "Unit price excl VAT",
            "VAT %",
            "Line amount excl VAT",
        ],
        [
            [
                invoice.source_filename,
                invoice.header.invoice_number,
                line.service_code,
                line.description,
                line.quantity,
                line.pricing_currency,
                line.exchange_rate,
                line.price_in_currency,
                line.currency_code,
                line.unit_price_ex_vat,
                line.vat_rate,
                line.line_amount_ex_vat,
            ]
            for invoice in invoices
            for line in invoice.lines
        ],
    )

    container_sheet = workbook.create_sheet("Containers")
    _write_rows(
        container_sheet,
        ["Source PDF", "Invoice number", "Container number", "Unit/type", "Weight", "CBM"],
        [
            [
                invoice.source_filename,
                invoice.header.invoice_number,
                container.container_number,
                container.unit_type,
                container.weight,
                container.cbm,
            ]
            for invoice in invoices
            for container in invoice.containers
        ],
    )

    atomic_save_workbook(workbook, output_path)


def write_validation_report(output_path: Path, outcomes: list[ValidationOutcome]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Validation"
    _write_rows(
        sheet,
        ["Source PDF", "Invoice number", "Status", "Errors", "Warnings"],
        [
            [
                outcome.source_pdf,
                outcome.invoice_number,
                "READY" if outcome.valid else "ERROR",
                "\n".join(outcome.errors),
                "\n".join(outcome.warnings),
            ]
            for outcome in outcomes
        ],
    )
    atomic_save_workbook(workbook, output_path)


def write_error_report(output_path: Path, outcomes: list[ValidationOutcome]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Errors"
    rows = []
    for outcome in outcomes:
        if outcome.valid:
            continue
        for error in outcome.errors:
            rows.append([outcome.source_pdf, outcome.invoice_number, error])
    _write_rows(sheet, ["Source PDF", "Invoice number", "Error"], rows)
    atomic_save_workbook(workbook, output_path)


def update_processed_log(log_path: Path, invoices: list[Invoice], outcomes: list[ValidationOutcome]) -> None:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    if log_path.exists():
        workbook = load_workbook(log_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Processed_Log"
        sheet.append(["Timestamp", "Filename", "Invoice number", "Client", "Total", "Status", "Message"])

    invoices_by_file = {invoice.source_filename: invoice for invoice in invoices}
    timestamp = iso_datetime_now()
    for outcome in outcomes:
        invoice = invoices_by_file.get(outcome.source_pdf)
        header = invoice.header if invoice else None
        sheet.append(
            [
                timestamp,
                outcome.source_pdf,
                outcome.invoice_number,
                header.client_name if header else "",
                header.total_amount_eur_incl_vat if header else "",
                "READY" if outcome.valid else "ERROR",
                "; ".join(outcome.warnings if outcome.valid else outcome.errors),
            ]
        )
    _style_sheet(sheet)
    atomic_save_workbook(workbook, log_path)


def _group_amounts_by_vat(invoice: Invoice) -> dict[Decimal, Decimal]:
    grouped: dict[Decimal, Decimal] = {}
    for line in invoice.lines:
        if line.vat_rate is None or line.line_amount_ex_vat is None:
            continue
        grouped[line.vat_rate] = grouped.get(line.vat_rate, Decimal("0")) + line.line_amount_ex_vat
    return grouped


def _find_header_row(sheet: Worksheet) -> int:
    for row_index in range(1, min(sheet.max_row, 100) + 1):
        row_values = [_normalize_header(sheet.cell(row=row_index, column=col).value) for col in range(1, sheet.max_column + 1)]
        required_matches = sum(1 for name in MONEO_COLUMNS if any(_header_matches(name, actual) for actual in row_values))
        if required_matches >= 10:
            return row_index
    raise TemplateError("Could not find the MONEO template header row. Set moneo.header_row in config.json.")


def _build_column_map(sheet: Worksheet, header_row: int) -> dict[str, int]:
    column_map: dict[str, int] = {}
    for required_name in MONEO_COLUMNS:
        for col in range(1, sheet.max_column + 1):
            actual = _normalize_header(sheet.cell(row=header_row, column=col).value)
            if _header_matches(required_name, actual):
                column_map[required_name] = col
                break

    missing = [name for name in MONEO_COLUMNS if name not in column_map]
    if missing:
        missing_text = ", ".join(missing)
        raise TemplateError(f"MONEO template is missing required columns: {missing_text}")
    return column_map


def _find_data_start_row(sheet: Worksheet, column_map: dict[str, int], header_row: int) -> int:
    candidate = header_row + 1
    non_formula_values = 0
    for name, col in column_map.items():
        if name in PROTECTED_FORMULA_COLUMNS:
            continue
        value = sheet.cell(row=candidate, column=col).value
        if value not in (None, ""):
            non_formula_values += 1
    if non_formula_values >= 3:
        candidate += 1
    return candidate


def _prepare_data_rows(sheet: Worksheet, column_map: dict[str, int], data_start_row: int, row_count: int) -> None:
    if row_count <= 0:
        return
    last_required_row = data_start_row + row_count - 1
    prototype_row = data_start_row
    max_col = max(column_map.values())
    for target_row in range(sheet.max_row + 1, last_required_row + 1):
        _copy_template_row(sheet, prototype_row, target_row, max_col)
    _clear_existing_data(sheet, column_map, data_start_row)


def _clear_existing_data(sheet: Worksheet, column_map: dict[str, int], data_start_row: int) -> None:
    if sheet.max_row < data_start_row:
        return
    for row in range(data_start_row, sheet.max_row + 1):
        for column_name, col in column_map.items():
            cell = sheet.cell(row=row, column=col)
            if column_name in PROTECTED_FORMULA_COLUMNS and _cell_has_formula(cell):
                continue
            cell.value = None


def _copy_template_row(sheet: Worksheet, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        source = sheet.cell(row=source_row, column=col)
        target = sheet.cell(row=target_row, column=col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if isinstance(source.value, str) and source.value.startswith("="):
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
        else:
            target.value = None


def _format_template_cell(cell, column_name: str) -> None:
    if column_name == "Rēķina datums":
        cell.number_format = "yyyy-mm-dd"
    elif column_name in {"Cena BEZ PVN", "Rindas summa BEZ PVN"}:
        cell.number_format = "#,##0.00"
    elif column_name == "Daudzums":
        cell.number_format = "#,##0"
    elif column_name == "Atlaides %":
        cell.number_format = "0.00"


def _write_rows(sheet: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet)


def _style_sheet(sheet: Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells[:50])
        width = min(max(max_len + 2, 12), 45)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _round_money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _header_matches(required_name: str, actual_normalized: str) -> bool:
    required = _normalize_header(required_name)
    return actual_normalized == required or actual_normalized.startswith(required)


def _cell_has_formula(cell) -> bool:
    return isinstance(cell.value, str) and cell.value.startswith("=")


def _validate_template_lists(workbook: Workbook, rows: list[MoneoRow], template_config: MoneoTemplateConfig) -> None:
    if "Set" not in workbook.sheetnames:
        return
    set_sheet = workbook["Set"]
    invoice_type_map = _read_lookup_table(set_sheet, key_col=1, value_col=2)
    vat_map = _read_lookup_table(set_sheet, key_col=8, value_col=9)

    expected_invoice_code = invoice_type_map.get(template_config.invoice_type)
    if expected_invoice_code and str(expected_invoice_code).strip() != template_config.invoice_type_value:
        raise TemplateError(
            f"Template maps invoice type {template_config.invoice_type!r} to {expected_invoice_code!r}, "
            f"but config uses {template_config.invoice_type_value!r}."
        )

    for row in rows:
        expected_pvn_code = vat_map.get(row.pvn_likme)
        if expected_pvn_code is None:
            raise TemplateError(f"PVN likme {row.pvn_likme!r} is not present in template Set sheet.")
        if str(expected_pvn_code).strip() != str(row.pvn_kods).strip():
            raise TemplateError(
                f"Template maps PVN likme {row.pvn_likme!r} to code {expected_pvn_code!r}, "
                f"but VAT mapping uses {row.pvn_kods!r}."
            )


def _read_lookup_table(sheet: Worksheet, key_col: int, value_col: int) -> dict[str, object]:
    values: dict[str, object] = {}
    for row_index in range(1, sheet.max_row + 1):
        key = sheet.cell(row=row_index, column=key_col).value
        value = sheet.cell(row=row_index, column=value_col).value
        if key in (None, "") or value in (None, ""):
            continue
        values[str(key).strip()] = value
    return values


def _force_excel_recalculation(workbook: Workbook) -> None:
    calculation = getattr(workbook, "calculation", None)
    if calculation is not None:
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
