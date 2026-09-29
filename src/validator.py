from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .exceptions import MappingError
from .invoice_models import Invoice, Money, ValidationOutcome
from .mapping_loader import MappingData


READY_STATUSES = {"READY", "IMPORTED"}


def load_processed_invoice_numbers(processed_log_path: Path) -> set[str]:
    if not processed_log_path.exists():
        return set()
    workbook = load_workbook(processed_log_path, data_only=True, read_only=True)
    sheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    try:
        invoice_col = headers.index("Invoice number") + 1
        status_col = headers.index("Status") + 1
    except ValueError:
        raise MappingError(
            f"Processed log is missing required columns 'Invoice number' and/or 'Status': {processed_log_path}"
        )

    processed: set[str] = set()
    for row in sheet.iter_rows(min_row=2):
        invoice_number = str(row[invoice_col - 1].value or "").strip()
        status = str(row[status_col - 1].value or "").strip().upper()
        if invoice_number and status in READY_STATUSES:
            processed.add(invoice_number)
    return processed


def ensure_processed_log(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Processed_Log"
    sheet.append(["Timestamp", "Filename", "Invoice number", "Client", "Total", "Status", "Message"])
    workbook.save(path)


def validate_invoice(
    invoice: Invoice,
    mappings: MappingData,
    processed_invoice_numbers: set[str],
    amount_tolerance: Money = Decimal("0.02"),
    template_available: bool = True,
    batch_invoice_numbers_seen: set[str] | None = None,
) -> ValidationOutcome:
    errors: list[str] = []
    warnings: list[str] = list(invoice.parser_warnings)
    header = invoice.header

    _require(header.invoice_number, "Invoice number is missing.", errors)
    _require(header.invoice_date, "Invoice date is missing.", errors)
    _require(header.client_vat_reg_number, "Client VAT/registration number is missing.", errors)
    _require(header.due_date, "Due date is missing.", errors)
    _require(header.payment_terms, "Payment terms are missing.", errors)
    _require(header.total_amount_eur_incl_vat, "Total amount including VAT is missing.", errors)

    if header.invoice_number and header.invoice_number in processed_invoice_numbers:
        errors.append(f"Duplicate invoice number already exists in processed log: {header.invoice_number}")
    if header.invoice_number and batch_invoice_numbers_seen and header.invoice_number in batch_invoice_numbers_seen:
        errors.append(f"Duplicate invoice number appears more than once in the current PDF batch: {header.invoice_number}")

    if header.invoice_date and header.due_date and header.due_date < header.invoice_date:
        errors.append(
            f"Due date ({header.due_date.isoformat()}) is before invoice date ({header.invoice_date.isoformat()})."
        )
    if header.invoice_date and header.invoice_date.year < 2000:
        errors.append(f"Invoice date is implausibly old: {header.invoice_date.isoformat()}.")
    if header.due_date and header.due_date.year < 2000:
        errors.append(f"Due date is implausibly old: {header.due_date.isoformat()}.")

    if not invoice.lines:
        errors.append("At least one invoice line is required.")
    if invoice.container_count == 0:
        errors.append("At least one container is required. Invoice will not be included in MONEO import.")
    if invoice.container_count < 0:
        errors.append(f"Container count is impossible: {invoice.container_count}.")

    client = mappings.find_client(header.client_name, header.client_vat_reg_number)
    if not header.client_name:
        if client is None:
            errors.append("Client name is missing and no client mapping could be matched by VAT/registration number.")
        else:
            warnings.append(
                "Client name was not extracted from the PDF; using exact client mapping matched by VAT/registration number."
            )
    if client is None:
        errors.append("Client was not found in Client_Mapping.xlsx.")

    vat_records = {}
    article_records = {}
    for vat_rate in sorted(_invoice_vat_rates(invoice)):
        article = mappings.find_article(vat_rate)
        if article is None:
            errors.append(f"No article mapping found for VAT rate {vat_rate}%.")
        else:
            article_records[vat_rate] = article

        vat_record = mappings.find_vat(vat_rate, client)
        if vat_record is None:
            errors.append(f"No VAT mapping found for VAT rate {vat_rate}% and this client.")
        elif not vat_record.pvn_likme or not vat_record.pvn_kods or vat_record.manual_review_required:
            errors.append(f"VAT rate {vat_rate}% requires manual review or missing PVN code in VAT_Mapping.xlsx.")
        else:
            vat_records[vat_rate] = vat_record

    for index, line in enumerate(invoice.lines, start=1):
        if line.vat_rate is None:
            errors.append(f"Line {index} has no VAT rate.")
        if line.line_amount_ex_vat is None:
            errors.append(f"Line {index} has no EUR amount excluding VAT.")
        elif line.line_amount_ex_vat < 0:
            errors.append(f"Line {index} has a negative amount excluding VAT: {line.line_amount_ex_vat}.")
        elif line.line_amount_ex_vat == 0:
            warnings.append(f"Line {index} has a zero amount excluding VAT.")
        if line.quantity is not None and line.quantity <= 0:
            errors.append(f"Line {index} has impossible quantity: {line.quantity}.")
        if line.unit_price_ex_vat is not None and line.unit_price_ex_vat < 0:
            errors.append(f"Line {index} has negative unit price excluding VAT: {line.unit_price_ex_vat}.")
        if line.vat_rate is not None and (line.vat_rate < 0 or line.vat_rate > 100):
            errors.append(f"Line {index} has impossible VAT rate: {line.vat_rate}%.")
        if line.currency_code and line.currency_code.upper() != "EUR":
            errors.append(f"Line {index} import currency code must be EUR, got {line.currency_code}.")
        if line.pricing_currency and line.pricing_currency.upper() != "EUR":
            if line.exchange_rate is None or line.price_in_currency is None:
                errors.append(
                    f"Line {index} has non-EUR pricing currency {line.pricing_currency} but missing exchange-rate context."
                )

    for index, container in enumerate(invoice.containers, start=1):
        if container.weight is not None and container.weight < 0:
            errors.append(f"Container {index} has negative weight: {container.container_number}.")
        if container.cbm is not None and container.cbm < 0:
            errors.append(f"Container {index} has negative CBM: {container.container_number}.")

    _validate_amounts(invoice, amount_tolerance, errors, warnings)

    if not template_available:
        errors.append("Official MONEO template is missing: templates/Rekinu_Imp_2024.xlsx.")

    return ValidationOutcome(
        source_pdf=invoice.source_filename,
        invoice_number=header.invoice_number,
        valid=not errors,
        errors=errors,
        warnings=warnings,
        client=client,
        vat_records_by_rate=vat_records,
        article_records_by_rate=article_records,
    )


def _invoice_vat_rates(invoice: Invoice) -> set[Money]:
    return {line.vat_rate for line in invoice.lines if line.vat_rate is not None}


def _validate_amounts(invoice: Invoice, tolerance: Money, errors: list[str], warnings: list[str]) -> None:
    line_sum = sum((line.line_amount_ex_vat or Decimal("0")) for line in invoice.lines)
    line_sum = _round_money(line_sum)
    header = invoice.header

    for label, value in [
        ("subtotal", header.subtotal),
        ("VAT amount", header.vat_amount),
        ("total including VAT", header.total_amount_eur_incl_vat),
    ]:
        if value is not None and value < 0:
            errors.append(f"Invoice {label} is negative: {value}.")

    if header.subtotal is not None and abs(line_sum - _round_money(header.subtotal)) > tolerance:
        errors.append(f"Line total excluding VAT ({line_sum}) does not match PDF subtotal ({header.subtotal}).")

    if (
        header.subtotal is not None
        and header.vat_amount is not None
        and header.total_amount_eur_incl_vat is not None
    ):
        expected_total = _round_money(header.subtotal + header.vat_amount)
        actual_total = _round_money(header.total_amount_eur_incl_vat)
        if abs(expected_total - actual_total) > tolerance:
            errors.append(
                f"Subtotal + VAT ({expected_total}) does not match total including VAT ({actual_total})."
            )

    if header.vat_amount is not None:
        expected_vat = sum(
            (line.line_amount_ex_vat or Decimal("0")) * ((line.vat_rate or Decimal("0")) / Decimal("100"))
            for line in invoice.lines
        )
        expected_vat = _round_money(expected_vat)
        if abs(expected_vat - _round_money(header.vat_amount)) > tolerance:
            errors.append(f"Calculated VAT ({expected_vat}) differs from PDF VAT amount ({header.vat_amount}).")


def _round_money(value: Money) -> Money:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _require(value: object, message: str, errors: list[str]) -> None:
    if value is None or value == "":
        errors.append(message)
